"""Creates PDF and Word versions of the resume (one column, standard fonts, no tables = ATS friendly)."""
from io import BytesIO
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate

from utils.ats import build_resume_text

HEADINGS = {
    "PROFESSIONAL SUMMARY", "SKILLS", "EXPERIENCE", "PROJECTS", "EDUCATION",
    "CERTIFICATIONS", "VOLUNTEER WORK", "INTERESTS & ADDITIONAL",
}


def _parse(text: str):
    lines = text.split("\n")
    name = lines[0].strip()
    contact = lines[1].strip() if len(lines) > 1 else ""
    sections = []
    for ln in lines[2:]:
        s = ln.strip()
        if not s:
            continue
        if s in HEADINGS:
            sections.append((s, []))
        elif sections:
            sections[-1][1].append(s)
    return name, contact, sections


def make_pdf(d: dict, role: str) -> bytes:
    name, contact, sections = _parse(build_resume_text(d, role))
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=1.8 * cm, rightMargin=1.8 * cm,
                            topMargin=1.5 * cm, bottomMargin=1.5 * cm,
                            title=f"{name} - Resume", author=name)
    s_name = ParagraphStyle("name", fontName="Helvetica-Bold", fontSize=18, alignment=TA_CENTER, spaceAfter=4)
    s_contact = ParagraphStyle("contact", fontName="Helvetica", fontSize=9.5, alignment=TA_CENTER, spaceAfter=8)
    s_head = ParagraphStyle("head", fontName="Helvetica-Bold", fontSize=11, spaceBefore=8, spaceAfter=2)
    s_body = ParagraphStyle("body", fontName="Helvetica", fontSize=10, leading=13)
    s_bullet = ParagraphStyle("bullet", parent=s_body, leftIndent=12, bulletIndent=2)

    story = [Paragraph(escape(name), s_name), Paragraph(escape(contact), s_contact)]
    for title, items in sections:
        story.append(Paragraph(escape(title), s_head))
        story.append(HRFlowable(width="100%", thickness=0.6, color=colors.black, spaceAfter=3))
        for it in items:
            if it.startswith("- "):
                story.append(Paragraph(escape(it[2:]), s_bullet, bulletText="•"))
            else:
                story.append(Paragraph(escape(it), s_body))
    doc.build(story)
    return buf.getvalue()


def make_docx(d: dict, role: str) -> bytes:
    name, contact, sections = _parse(build_resume_text(d, role))
    doc = Document()
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(1.8)
        sec.top_margin = sec.bottom_margin = Cm(1.5)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(name)
    r.bold = True
    r.font.size = Pt(18)
    p = doc.add_paragraph(contact)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for title, items in sections:
        h = doc.add_paragraph()
        run = h.add_run(title)
        run.bold = True
        run.font.size = Pt(11.5)
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(2)
        for it in items:
            if it.startswith("- "):
                doc.add_paragraph(it[2:], style="List Bullet")
            else:
                doc.add_paragraph(it)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()